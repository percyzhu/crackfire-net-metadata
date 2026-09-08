from pathlib import Path
import hashlib,json
import fitz
import openpyxl

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
PDF=ROOT/'3_最终成果/论文/LaTeX源码/cn/figures/mesh_sensitivity.pdf'
XLSX=ROOT/'1_代码/src/data/T-t_experimental_curve.xlsx'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
doc=fitz.open(PDF);page=doc[0]
page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(HERE/'original_mesh_validation.png')
(HERE/'original_pdf_text.txt').write_text(page.get_text(),encoding='utf-8')
drawings=[]
for i,d in enumerate(page.get_drawings()):
    drawings.append({'index':i,'color':d['color'],'fill':d['fill'],'width':d['width'],
        'dashes':d['dashes'],'rect':list(d['rect']),'items':len(d['items']),
        'item_types':sorted(set(k[0] for k in d['items'])),
        'first_items':[str(k) for k in d['items'][:2]]})
(HERE/'original_pdf_drawing_inventory.json').write_text(json.dumps({'source':str(PDF),'sha256':sha(PDF),'page_rect':list(page.rect),'drawings':drawings},indent=2),encoding='utf-8')
wb=openpyxl.load_workbook(XLSX,read_only=True,data_only=True)
reports={}
for s in wb:
    rows=list(s.iter_rows(values_only=True))
    reports[s.title]={'rows':s.max_row,'cols':s.max_column,'headers':list(rows[0]),
        'numeric_counts':{col:sum(len(row)>j and isinstance(row[j],(int,float)) for row in rows[1:])
                          for j,col in zip(range(10,14),'KLMN')} if s.title!='Sheet1' else None}
(HERE/'source_workbook_headers.json').write_text(json.dumps({'source':str(XLSX),'sha256':sha(XLSX),'sheets':reports},indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'pdf_size':list(page.rect),'text':page.get_text(),'colored_polyline_candidates':[d for d in drawings if d['color'] and max(d['color'])-min(d['color'])>.02 and d['items']>20]},indent=2,ensure_ascii=False))
